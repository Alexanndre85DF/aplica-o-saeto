from __future__ import annotations

from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .cadastro import _eh_placeholder, chave_municipio
from .regras import periodo_dias

VALOR_PADRAO = Decimal("189")
SEDE = "GURUPI"
META_VALOR = "diaria_valor_padrao"
CENTAVO = Decimal("0.01")
MEIA = Decimal("0.5")


def _dec(valor, padrao: Decimal | None = None) -> Decimal:
    if valor is None or str(valor).strip() == "":
        if padrao is None:
            raise ValueError("Informe um valor.")
        return padrao
    texto = str(valor).strip().replace("R$", "").replace(" ", "")
    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".")
    try:
        return Decimal(texto)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("Valor inválido.") from exc


def _qtd(valor, padrao: Decimal) -> Decimal:
    qtd = _dec(valor, padrao)
    if qtd <= 0:
        raise ValueError("A quantidade de diárias deve ser maior que zero.")
    meios = (qtd / MEIA).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    return (meios * MEIA).quantize(MEIA)


def _dinheiro(valor) -> Decimal:
    return _dec(valor, VALOR_PADRAO).quantize(CENTAVO)


def fmt_moeda(valor: Decimal | float | str | None) -> str:
    q = _dinheiro(valor if valor is not None else 0)
    texto = f"{q:.2f}"
    reais, cents = texto.split(".")
    milhar = ""
    while len(reais) > 3:
        milhar = "." + reais[-3:] + milhar
        reais = reais[:-3]
    return f"R$ {reais}{milhar},{cents}"


def fmt_periodo(saida: str | None, retorno: str | None) -> str:
    if not saida and not retorno:
        return "—"
    ini = (saida or retorno)[:10]
    fim = (retorno or saida)[:10]

    def curta(iso: str) -> str:
        _ano, mes, dia = iso.split("-")
        return f"{dia}/{mes}"

    return f"{curta(ini)} A {curta(fim)}"


def _eh_sede(nome: str | None) -> bool:
    chave = chave_municipio(nome or "")
    return chave == SEDE or chave.startswith(SEDE + " ")


def _meta(conn, chave: str, padrao: str | None = None) -> str | None:
    row = conn.execute("SELECT valor FROM meta WHERE chave = ?", (chave,)).fetchone()
    if not row or row["valor"] in (None, ""):
        return padrao
    return row["valor"]


def _gravar_meta(conn, chave: str, valor: str) -> None:
    conn.execute(
        "INSERT OR REPLACE INTO meta(chave, valor) VALUES (?, ?)",
        (chave, valor),
    )


def valor_padrao(conn) -> Decimal:
    return _dinheiro(_meta(conn, META_VALOR, str(VALOR_PADRAO)))


def _valores_por_municipio(conn) -> dict[int, Decimal]:
    try:
        rows = conn.execute("SELECT municipio_id, valor FROM diaria_valores").fetchall()
    except Exception:
        return {}
    return {int(r["municipio_id"]): _dinheiro(r["valor"]) for r in rows}


def _ajustes(conn) -> dict[tuple[int, int], dict]:
    try:
        rows = conn.execute("SELECT * FROM diaria_ajustes").fetchall()
    except Exception:
        return {}
    return {(int(r["aplicador_id"]), int(r["municipio_id"])): dict(r) for r in rows}


def _viagens(conn) -> dict[int, dict]:
    rows = conn.execute(
        "SELECT municipio_id, data_saida, data_retorno FROM viagens"
    ).fetchall()
    return {int(r["municipio_id"]): dict(r) for r in rows}


def _ocupacoes(conn) -> list[dict]:
    titulares = conn.execute(
        """SELECT v.aplicador_id, e.municipio_id, m.nome AS municipio,
                  v.data, 'titular' AS papel
           FROM vagas v
           JOIN escolas e ON e.id = v.escola_id
           JOIN municipios m ON m.id = e.municipio_id
           WHERE v.aplicador_id IS NOT NULL"""
    ).fetchall()
    extras = []
    try:
        extras = conn.execute(
            """SELECT DISTINCT x.aplicador_id, e.municipio_id, m.nome AS municipio,
                      v.data, 'extra' AS papel
               FROM vaga_extras x
               JOIN vagas v ON v.id = x.vaga_id
               JOIN escolas e ON e.id = v.escola_id
               JOIN municipios m ON m.id = e.municipio_id"""
        ).fetchall()
    except Exception:
        extras = []
    return [dict(r) for r in list(titulares) + list(extras)]


def _periodo_municipio(viagem: dict | None, datas: list[str]) -> tuple[str | None, str | None]:
    if viagem:
        saida = viagem.get("data_saida")
        retorno = viagem.get("data_retorno") or saida
        if saida:
            return saida[:10], (retorno or saida)[:10]
    validas = sorted({d[:10] for d in datas if d})
    if not validas:
        return None, None
    return validas[0], validas[-1]


def _qtd_sugerida(saida: str | None, retorno: str | None) -> Decimal:
    dias = periodo_dias(saida, retorno)
    if dias:
        return Decimal(len(dias))
    return Decimal("1")


def montar_folha(conn) -> dict:
    padrao = valor_padrao(conn)
    valores = _valores_por_municipio(conn)
    ajustes = _ajustes(conn)
    viagens = _viagens(conn)
    pessoas = {
        int(r["id"]): dict(r)
        for r in conn.execute("SELECT * FROM aplicadores").fetchall()
    }

    agrupado: dict[tuple[int, int], dict] = {}
    for ocup in _ocupacoes(conn):
        apl_id = int(ocup["aplicador_id"])
        mun_id = int(ocup["municipio_id"])
        pessoa = pessoas.get(apl_id)
        if not pessoa:
            continue
        if _eh_placeholder(pessoa.get("nome"), pessoa.get("codigo")):
            continue
        if _eh_sede(ocup.get("municipio")):
            continue
        chave = (apl_id, mun_id)
        item = agrupado.setdefault(
            chave,
            {
                "aplicador_id": apl_id,
                "municipio_id": mun_id,
                "municipio": ocup["municipio"],
                "papeis": set(),
                "datas": [],
            },
        )
        item["papeis"].add(ocup["papel"])
        if ocup.get("data"):
            item["datas"].append(ocup["data"])

    grupos_map: dict[tuple, dict] = {}
    linhas_avulsas = []
    for (apl_id, mun_id), bruto in agrupado.items():
        pessoa = pessoas[apl_id]
        saida, retorno = _periodo_municipio(viagens.get(mun_id), bruto["datas"])
        sugerida = _qtd_sugerida(saida, retorno)
        ajuste = ajustes.get((apl_id, mun_id), {})
        excluido = bool(ajuste.get("excluido"))
        qtd = (
            _qtd(ajuste.get("qtd_diarias"), sugerida)
            if ajuste.get("qtd_diarias") not in (None, "")
            else sugerida
        )
        unitario = valores.get(mun_id, padrao)
        valor = (qtd * unitario).quantize(CENTAVO)
        papeis = sorted(bruto["papeis"])
        linha = {
            "aplicador_id": apl_id,
            "municipio_id": mun_id,
            "nome": pessoa.get("nome") or pessoa.get("codigo"),
            "codigo": pessoa.get("codigo"),
            "tipo": pessoa.get("tipo") or "APLICADOR",
            "matricula": (pessoa.get("matricula") or "").strip(),
            "papeis": papeis,
            "papel": "ambos" if len(papeis) > 1 else papeis[0],
            "os": (ajuste.get("os") or "").strip(),
            "qtd_diarias": float(qtd),
            "qtd_sugerida": float(sugerida),
            "qtd_ajustada": ajuste.get("qtd_diarias") not in (None, ""),
            "valor_unitario": float(unitario),
            "valor": float(valor),
            "valor_fmt": fmt_moeda(valor),
            "excluido": excluido,
            "sem_matricula": not (pessoa.get("matricula") or "").strip(),
            "data_saida": saida,
            "data_retorno": retorno,
            "data_fmt": fmt_periodo(saida, retorno),
            "municipio": bruto["municipio"],
            "rota": f"{SEDE} A {str(bruto['municipio'] or '').upper()}",
        }
        linhas_avulsas.append(linha)
        gchave = (saida or "", retorno or "", mun_id)
        grupo = grupos_map.setdefault(
            gchave,
            {
                "data_saida": saida,
                "data_retorno": retorno,
                "data_fmt": fmt_periodo(saida, retorno),
                "municipio_id": mun_id,
                "municipio": bruto["municipio"],
                "rota": linha["rota"],
                "linhas": [],
            },
        )
        grupo["linhas"].append(linha)

    grupos = []
    total_geral = Decimal("0")
    pessoas_ids = set()
    sem_mat = 0
    excluidas = 0
    for chave in sorted(grupos_map, key=lambda k: (k[0] or "9999", k[2], k[1] or "")):
        grupo = grupos_map[chave]
        grupo["linhas"].sort(key=lambda x: (x["nome"] or "").upper())
        soma = Decimal("0")
        for linha in grupo["linhas"]:
            if linha["excluido"]:
                excluidas += 1
                continue
            soma += Decimal(str(linha["valor"]))
            pessoas_ids.add(linha["aplicador_id"])
            if linha["sem_matricula"]:
                sem_mat += 1
        grupo["total"] = float(soma)
        grupo["total_fmt"] = fmt_moeda(soma)
        grupo["pessoas"] = sum(1 for l in grupo["linhas"] if not l["excluido"])
        grupos.append(grupo)
        total_geral += soma

    destinos = []
    vistos = set()
    for grupo in grupos:
        mid = grupo["municipio_id"]
        if mid in vistos:
            continue
        vistos.add(mid)
        destinos.append(
            {
                "municipio_id": mid,
                "municipio": grupo["municipio"],
                "valor": float(valores.get(mid, padrao)),
                "valor_fmt": fmt_moeda(valores.get(mid, padrao)),
                "customizado": mid in valores,
            }
        )
    destinos.sort(key=lambda d: d["municipio"])

    return {
        "valor_padrao": float(padrao),
        "valor_padrao_fmt": fmt_moeda(padrao),
        "sede": SEDE,
        "valores": destinos,
        "grupos": grupos,
        "linhas": sorted(
            linhas_avulsas,
            key=lambda x: (x["data_saida"] or "9999", x["municipio"], x["nome"] or ""),
        ),
        "totais": {
            "pessoas": len(pessoas_ids),
            "linhas": sum(g["pessoas"] for g in grupos),
            "viagens": len(grupos),
            "excluidas": excluidas,
            "sem_matricula": sem_mat,
            "valor": float(total_geral),
            "valor_fmt": fmt_moeda(total_geral),
        },
    }


def gravar_config(conn, valor_padrao_novo=None, valores=None) -> dict:
    if valor_padrao_novo is not None:
        _gravar_meta(conn, META_VALOR, str(_dinheiro(valor_padrao_novo)))
    if valores is not None:
        for item in valores:
            mid = int(item["municipio_id"])
            mun = conn.execute("SELECT id FROM municipios WHERE id = ?", (mid,)).fetchone()
            if not mun:
                raise LookupError("Município não encontrado.")
            if item.get("valor") in (None, ""):
                conn.execute("DELETE FROM diaria_valores WHERE municipio_id = ?", (mid,))
                continue
            valor = str(_dinheiro(item["valor"]))
            existe = conn.execute(
                "SELECT 1 FROM diaria_valores WHERE municipio_id = ?", (mid,)
            ).fetchone()
            if existe:
                conn.execute(
                    "UPDATE diaria_valores SET valor = ? WHERE municipio_id = ?",
                    (valor, mid),
                )
            else:
                conn.execute(
                    "INSERT INTO diaria_valores(municipio_id, valor) VALUES (?, ?)",
                    (mid, valor),
                )
    return montar_folha(conn)


def gravar_ajuste(
    conn,
    aplicador_id: int,
    municipio_id: int,
    os=None,
    qtd_diarias=None,
    excluido=None,
) -> dict:
    apl = conn.execute("SELECT id FROM aplicadores WHERE id = ?", (aplicador_id,)).fetchone()
    if not apl:
        raise LookupError("Aplicador não encontrado.")
    mun = conn.execute("SELECT id FROM municipios WHERE id = ?", (municipio_id,)).fetchone()
    if not mun:
        raise LookupError("Município não encontrado.")
    atual = conn.execute(
        """SELECT * FROM diaria_ajustes
           WHERE aplicador_id = ? AND municipio_id = ?""",
        (aplicador_id, municipio_id),
    ).fetchone()
    os_val = atual["os"] if atual and os is None else os
    qtd_val = atual["qtd_diarias"] if atual and qtd_diarias is None else qtd_diarias
    if qtd_diarias == "":
        qtd_val = None
    elif qtd_diarias is not None:
        qtd_val = str(_qtd(qtd_diarias, Decimal("1")))
    excl = atual["excluido"] if atual and excluido is None else (1 if excluido else 0)
    os_val = (os_val or "").strip() or None
    if not os_val and qtd_val is None and not excl:
        if atual:
            conn.execute("DELETE FROM diaria_ajustes WHERE id = ?", (atual["id"],))
        return montar_folha(conn)
    if atual:
        conn.execute(
            """UPDATE diaria_ajustes
               SET os = ?, qtd_diarias = ?, excluido = ?
               WHERE id = ?""",
            (os_val, qtd_val, 1 if excl else 0, atual["id"]),
        )
    else:
        conn.execute(
            """INSERT INTO diaria_ajustes(aplicador_id, municipio_id, os, qtd_diarias, excluido)
               VALUES (?, ?, ?, ?, ?)""",
            (aplicador_id, municipio_id, os_val, qtd_val, 1 if excl else 0),
        )
    return montar_folha(conn)


def exportar_xlsx(conn) -> tuple[BytesIO, str]:
    folha = montar_folha(conn)
    wb = Workbook()
    ws = wb.active
    ws.title = "Diárias"
    cab = ["DATA", "MUNICÍPIO", "SERVIDOR", "MATRÍCULA", "OS", "VALOR"]
    borda = Border(
        left=Side(style="thin", color="808080"),
        right=Side(style="thin", color="808080"),
        top=Side(style="thin", color="808080"),
        bottom=Side(style="thin", color="808080"),
    )
    fundo_cab = PatternFill("solid", fgColor="D9E2F3")
    fundo_tot = PatternFill("solid", fgColor="FFF2CC")
    negrito = Font(bold=True)
    ws.append(cab)
    for cell in ws[1]:
        cell.font = negrito
        cell.fill = fundo_cab
        cell.border = borda
        cell.alignment = Alignment(horizontal="center")
    for grupo in folha["grupos"]:
        for linha in grupo["linhas"]:
            if linha["excluido"]:
                continue
            ws.append(
                [
                    linha["data_fmt"],
                    linha["rota"],
                    (linha["nome"] or "").upper(),
                    linha["matricula"],
                    linha["os"],
                    linha["valor_fmt"],
                ]
            )
            for cell in ws[ws.max_row]:
                cell.border = borda
        ws.append(["", "", f"Total: {grupo['total_fmt']}", "", "", f"Total: {grupo['total_fmt']}"])
        for cell in ws[ws.max_row]:
            cell.font = negrito
            cell.fill = fundo_tot
            cell.border = borda
    ws.append([])
    ws.append(["", "", f"Total geral: {folha['totais']['valor_fmt']}", "", "", folha["totais"]["valor_fmt"]])
    for cell in ws[ws.max_row]:
        cell.font = negrito
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 18
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio, "diarias-saeto.xlsx"
