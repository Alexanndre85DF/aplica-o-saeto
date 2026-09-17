from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .cadastro import consolidar_municipios, garantir_municipio
from .regras import (
    eh_segundo_ano_dia1,
    normalizar_serie,
    normalizar_texto,
    normalizar_turno,
    serie_par_segundo_ano,
)

REGIONAL_PADRAO = "GURUPI"


def eh_planilha_homologacao(wb) -> bool:
    nomes = {str(n).strip().upper() for n in wb.sheetnames}
    return "ALUNOS POR TURMA" in nomes or any(
        "TURMA" in n and "TURNO" in n for n in nomes
    )


def _cab(valor) -> str:
    return normalizar_texto(valor).upper()


def _indice(mapa: dict[str, int], *chaves: str) -> int | None:
    for chave in chaves:
        if chave in mapa:
            return mapa[chave]
    for cab, idx in mapa.items():
        for chave in chaves:
            if chave in cab:
                return idx
    return None


def _mapa_cabecalho(ws, linha: int = 2) -> dict[str, int]:
    mapa = {}
    for c in range(1, (ws.max_column or 1) + 1):
        h = _cab(ws.cell(linha, c).value)
        if h:
            mapa[h] = c
    return mapa


def _valor(row, idx: int | None):
    if not idx or idx < 1 or idx > len(row):
        return None
    return row[idx - 1]


def serie_da_etapa(texto: str) -> str | None:
    t = _cab(texto)
    t = t.replace("SERIE", "SÉRIE").replace("SERÍE", "SÉRIE")
    partes = [t]
    if "-" in t:
        partes.append(t.split("-")[0].strip())
    for parte in partes:
        if "SÉRIE" in parte and "3" in parte:
            return normalizar_serie("3ª SÉRIE")
        if "SÉRIE" in parte and "2" in parte:
            return normalizar_serie("2ª SÉRIE")
        if "ANO" in parte and "9" in parte:
            return normalizar_serie("9º ANO")
        if "ANO" in parte and "8" in parte:
            return normalizar_serie("8º ANO")
        if "ANO" in parte and "5" in parte:
            return normalizar_serie("5º ANO")
        if "ANO" in parte and "4" in parte:
            return normalizar_serie("4º ANO")
        if "ANO" in parte and "2" in parte:
            return normalizar_serie("2º ANO - DIA 1")
    return None


def _rede_de(texto) -> str:
    rede = _cab(texto)
    if "MUNICIPAL" in rede:
        return "MUNICIPAL"
    if "CONVEN" in rede or "PRIVAD" in rede:
        return "CONVENIADA"
    return "ESTADUAL"


def eh_planilha_caed(wb) -> bool:
    nomes = {str(n).strip().upper() for n in wb.sheetnames}
    if "CAED" in nomes:
        return True
    for nome in wb.sheetnames:
        ws = wb[nome]
        headers = {
            _cab(ws.cell(1, c).value)
            for c in range(1, min(40, (ws.max_column or 1) + 1))
        }
        if "NM_ALUNO" in headers and ("NM_TURMA" in headers or "NM_ESCOLA" in headers):
            return True
    return False


def _aba_caed(wb):
    for nome in wb.sheetnames:
        if str(nome).strip().upper() == "CAED":
            return wb[nome]
    for nome in wb.sheetnames:
        ws = wb[nome]
        headers = {
            _cab(ws.cell(1, c).value)
            for c in range(1, min(40, (ws.max_column or 1) + 1))
        }
        if "NM_ALUNO" in headers:
            return ws
    return None


def importar_caed(caminho: Path, conn, regional: str = REGIONAL_PADRAO) -> dict:
    wb = load_workbook(caminho, data_only=True)
    ws = _aba_caed(wb)
    if ws is None:
        raise ValueError("Esta planilha não tem a aba Caed de alunos.")
    mapa = _mapa_cabecalho(ws, 1)
    col_reg = _indice(mapa, "NM_REGIONAL", "REGIONAL")
    col_mun = _indice(mapa, "NM_MUNICIPIO", "NM_MUNICÍPIO", "MUNICIPIO")
    col_esc = _indice(mapa, "NM_ESCOLA", "ESCOLA")
    col_cod = _indice(mapa, "CD_ESCOLA")
    col_rede = _indice(mapa, "DC_REDE_ENSINO", "DC_REDE", "REDE")
    col_loc = _indice(mapa, "DC_LOCALIZACAO", "DC_LOCALIZACAO_ESCOLA", "LOCALIZACAO")
    col_turma = _indice(mapa, "NM_TURMA")
    col_turno = _indice(mapa, "DC_TURNO", "DC_TURNO_TURMA")
    col_etapa = _indice(mapa, "DC_ETAPA", "DC_ETAPA_DIVULGACAO_TURMA", "DC_ETAPA_TURMA")
    if not col_esc or not col_mun:
        raise ValueError("Não achei escola e município na planilha Caed.")

    alvo = _cab(regional)
    avisos = []
    lidas = 0
    ignoradas = 0
    grupos: dict[tuple, dict] = {}

    for row in ws.iter_rows(min_row=2, max_col=40, values_only=True):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        lidas += 1
        regional_row = _cab(_valor(row, col_reg))
        if alvo and regional_row and alvo not in regional_row:
            ignoradas += 1
            continue
        mun = _cab(_valor(row, col_mun))
        escola = _cab(_valor(row, col_esc))
        if not mun or not escola:
            continue
        turma = normalizar_texto(_valor(row, col_turma))
        serie = serie_da_etapa(str(_valor(row, col_etapa) or "")) or serie_da_etapa(turma)
        if not serie:
            avisos.append(f"{escola}: turma {turma or '?'} sem série reconhecida.")
            continue
        turno = normalizar_turno(_valor(row, col_turno))
        codigo = str(_valor(row, col_cod) or "").strip() or f"CAD-{escola[:20]}"
        chave = (codigo, turma, serie, turno)
        if chave not in grupos:
            grupos[chave] = {
                "mun": mun,
                "escola": escola,
                "codigo": codigo,
                "rede": _rede_de(_valor(row, col_rede)),
                "rural": 1 if "RURAL" in _cab(_valor(row, col_loc)) else 0,
                "turma": turma,
                "serie": serie,
                "turno": turno,
                "n_alunos": 0,
            }
        grupos[chave]["n_alunos"] += 1

    ordem_por_chave: dict[tuple, int] = {}
    for r in conn.execute(
        "SELECT escola_id, serie, turno, COALESCE(MAX(ordem), 0) n FROM vagas GROUP BY escola_id, serie, turno"
    ):
        ordem_por_chave[(r["escola_id"], r["serie"], r["turno"])] = r["n"]

    gravadas = 0
    ja_tinham = 0
    for item in grupos.values():
        municipio_id = garantir_municipio(conn, item["mun"])
        existente = conn.execute(
            "SELECT id FROM escolas WHERE codigo = ?", (item["codigo"],)
        ).fetchone()
        if existente:
            escola_id = existente[0]
            conn.execute(
                "UPDATE escolas SET nome=?, municipio_id=?, rede=?, rural=? WHERE id=?",
                (item["escola"], municipio_id, item["rede"], item["rural"], escola_id),
            )
        else:
            cur = conn.execute(
                """INSERT INTO escolas(codigo, nome, municipio_id, rede, rural)
                   VALUES (?, ?, ?, ?, ?)""",
                (item["codigo"], item["escola"], municipio_id, item["rede"], item["rural"]),
            )
            escola_id = cur.lastrowid
        ja = conn.execute(
            """SELECT 1 FROM vagas
               WHERE escola_id = ? AND serie = ? AND turno = ?
                 AND COALESCE(turma, '') = ?""",
            (escola_id, item["serie"], item["turno"], item["turma"] or ""),
        ).fetchone()
        if ja:
            ja_tinham += 1
            continue
        chave_ord = (escola_id, item["serie"], item["turno"])
        ordem_por_chave[chave_ord] = ordem_por_chave.get(chave_ord, 0) + 1
        conn.execute(
            """INSERT INTO vagas(escola_id, serie, turno, data, ordem, origem_linha, status, turma, n_alunos)
               VALUES (?, ?, ?, ?, ?, ?, 'PREVISTA', ?, ?)""",
            (
                escola_id,
                item["serie"],
                item["turno"],
                None,
                ordem_por_chave[chave_ord],
                None,
                item["turma"],
                item["n_alunos"],
            ),
        )
        gravadas += 1

    consolidar_municipios(conn)
    return {
        "arquivo": caminho.name,
        "linhas_lidas": lidas,
        "turmas": len(grupos),
        "vagas_gravadas": gravadas,
        "ja_existiam": ja_tinham,
        "ignoradas_fora_regional": ignoradas,
        "avisos": avisos[:40],
        "regional": regional,
        "acrescentar": True,
    }


def _int(valor) -> int | None:
    if valor is None or str(valor).strip() in {"", "-"}:
        return None
    try:
        return int(float(valor))
    except (TypeError, ValueError):
        return None


def importar_homologacao(
    caminho: Path, conn, regional: str = REGIONAL_PADRAO, so_faltantes: bool = False
) -> dict:
    wb = load_workbook(caminho, data_only=True)
    if "Alunos por turma" not in wb.sheetnames:
        raise ValueError("Esta planilha não tem a aba Alunos por turma.")
    ws = wb["Alunos por turma"]
    mapa = _mapa_cabecalho(ws, 2)
    col_reg = _indice(mapa, "NM_REGIONAL", "REGIONAL")
    col_mun = _indice(mapa, "NM_MUNICIPIO", "NM_MUNICÍPIO", "MUNICIPIO")
    col_esc = _indice(mapa, "NM_ESCOLA", "ESCOLA")
    col_cod = _indice(mapa, "CD_ESCOLA")
    col_rede = _indice(mapa, "DC_REDE", "REDE")
    col_loc = _indice(mapa, "DC_LOCALIZACAO_ESCOLA", "LOCALIZACAO")
    col_turma = _indice(mapa, "NM_TURMA")
    col_cd_turma = _indice(mapa, "CD_TURMA")
    col_turno = _indice(mapa, "DC_TURNO_TURMA", "DC_TURNO")
    col_etapa = _indice(mapa, "DC_ETAPA_DIVULGACAO_TURMA", "DC_ETAPA_TURMA")
    col_alunos = _indice(mapa, "QT_ALUNO")
    if not col_esc or not col_mun:
        raise ValueError("Não achei escola e município na aba Alunos por turma.")

    alvo = _cab(regional)
    avisos = []
    lidas = 0
    gravadas = 0
    ignoradas = 0
    if not so_faltantes:
        conn.execute("DELETE FROM vagas WHERE origem_linha IS NOT NULL")
    ordem_por_chave: dict[tuple, int] = {}
    for r in conn.execute("SELECT escola_id, serie, turno, COALESCE(MAX(ordem), 0) n FROM vagas GROUP BY escola_id, serie, turno"):
        ordem_por_chave[(r["escola_id"], r["serie"], r["turno"])] = r["n"]

    for row in ws.iter_rows(min_row=3, max_col=32, values_only=True):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        lidas += 1
        regional_row = _cab(_valor(row, col_reg))
        if alvo and regional_row and alvo not in regional_row:
            ignoradas += 1
            continue
        mun = _cab(_valor(row, col_mun))
        escola = _cab(_valor(row, col_esc))
        if not mun or not escola:
            avisos.append(f"Linha sem município ou escola, ignorada.")
            continue
        codigo = str(_valor(row, col_cod) or "").strip() or f"CAD-{escola[:20]}"
        rede = _rede_de(_valor(row, col_rede))
        rural = 1 if "RURAL" in _cab(_valor(row, col_loc)) else 0
        turma = normalizar_texto(_valor(row, col_turma))
        serie = serie_da_etapa(str(_valor(row, col_etapa) or "")) or serie_da_etapa(turma)
        if not serie:
            avisos.append(f"{escola}: turma {turma or '?'} sem série reconhecida.")
            continue
        turno = normalizar_turno(_valor(row, col_turno))
        n_alunos = _int(_valor(row, col_alunos))
        cd_turma = _int(_valor(row, col_cd_turma))

        municipio_id = garantir_municipio(conn, mun)

        existente = conn.execute(
            "SELECT id FROM escolas WHERE codigo = ?", (codigo,)
        ).fetchone()
        if existente:
            escola_id = existente[0]
            conn.execute(
                "UPDATE escolas SET nome=?, municipio_id=?, rede=?, rural=? WHERE id=?",
                (escola, municipio_id, rede, rural, escola_id),
            )
        else:
            cur = conn.execute(
                """INSERT INTO escolas(codigo, nome, municipio_id, rede, rural)
                   VALUES (?, ?, ?, ?, ?)""",
                (codigo, escola, municipio_id, rede, rural),
            )
            escola_id = cur.lastrowid

        origem = cd_turma
        if so_faltantes and origem:
            ja = conn.execute(
                "SELECT 1 FROM vagas WHERE origem_linha = ?", (origem,)
            ).fetchone()
            if ja:
                continue
        chave = (escola_id, serie, turno)
        ordem_por_chave[chave] = ordem_por_chave.get(chave, 0) + 1
        ordem = ordem_por_chave[chave]
        conn.execute(
            """INSERT INTO vagas(escola_id, serie, turno, data, ordem, origem_linha, status, turma, n_alunos)
               VALUES (?, ?, ?, ?, ?, ?, 'PREVISTA', ?, ?)""",
            (escola_id, serie, turno, None, ordem, origem, turma, n_alunos),
        )
        gravadas += 1
        if eh_segundo_ano_dia1(serie):
            par = serie_par_segundo_ano(serie)
            origem_par = -origem if origem else None
            if so_faltantes and origem_par:
                ja_par = conn.execute(
                    "SELECT 1 FROM vagas WHERE origem_linha = ?", (origem_par,)
                ).fetchone()
                if ja_par:
                    continue
            chave_par = (escola_id, par, turno)
            ordem_por_chave[chave_par] = ordem_por_chave.get(chave_par, 0) + 1
            conn.execute(
                """INSERT INTO vagas(escola_id, serie, turno, data, ordem, origem_linha, status, turma, n_alunos)
                   VALUES (?, ?, ?, ?, ?, ?, 'PREVISTA', ?, ?)""",
                (
                    escola_id,
                    par,
                    turno,
                    None,
                    ordem_por_chave[chave_par],
                    origem_par,
                    turma,
                    n_alunos,
                ),
            )
            gravadas += 1

    consolidar_municipios(conn)
    return {
        "arquivo": caminho.name,
        "linhas_lidas": lidas,
        "vagas_gravadas": gravadas,
        "ignoradas_fora_regional": ignoradas,
        "avisos": avisos[:40],
        "regional": regional,
    }
