from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .cadastro import consolidar_municipios, garantir_municipio, id_municipio_por_nome
from .config import DATA_DIR
from .db import get_db
from .regras import (
    eh_rural,
    inferir_rede,
    iso,
    normalizar_serie,
    normalizar_texto,
    normalizar_turno,
    parse_data,
)

COLUNAS_PADRAO = {
    "mun": 1,
    "codigo": 2,
    "escola": 3,
    "serie": 4,
    "turno": 5,
    "data": 6,
    "aplicador": 7,
    "saida": 8,
    "retorno": 9,
}


def _cab(valor) -> str:
    return normalizar_texto(valor).upper()


def _mapear_colunas(ws) -> tuple[int, dict[str, int]]:
    for r in range(1, 8):
        headers = [_cab(ws.cell(r, c).value) for c in range(1, 16)]
        texto = " ".join(headers)
        if "MUNICIPIO" not in texto and "MUNICÍPIO" not in texto:
            continue
        if "ESCOLA" not in texto and "TURNO" not in texto and "APLICADOR" not in texto:
            continue
        mapa: dict[str, int] = {}
        for i, h in enumerate(headers, start=1):
            if not h:
                continue
            if "MUNICIPIO" in h or "MUNICÍPIO" in h:
                mapa["mun"] = i
            elif "INEP" in h or h in {"CD_ESCOLA", "CODIGO", "CÓDIGO", "CODIGO ESCOLA"}:
                mapa["codigo"] = i
            elif "ESCOLA" in h:
                mapa.setdefault("escola", i)
            elif "TURNO" in h:
                mapa["turno"] = i
            elif "APLICADOR" in h:
                mapa["aplicador"] = i
            elif "SAI" in h:
                mapa["saida"] = i
            elif "RETOR" in h:
                mapa["retorno"] = i
            elif "SERIE" in h or "SÉRIE" in h or h in {"ANO", "ETAPA", "SERIE/ANO"}:
                mapa["serie"] = i
            elif h.startswith("DATA"):
                mapa.setdefault("data", i)
        if "serie" not in mapa and "escola" in mapa and "turno" in mapa:
            for c in range(mapa["escola"] + 1, mapa["turno"]):
                if c not in mapa.values():
                    mapa["serie"] = c
                    break
        if "mun" in mapa and "escola" in mapa:
            completo = {**COLUNAS_PADRAO, **mapa}
            return r, completo
    return 2, dict(COLUNAS_PADRAO)


def _valor(row, indice: int):
    if not indice or indice < 1 or indice > len(row):
        return None
    return row[indice - 1].value


def _codigo_escola(codigo: str, escola: str) -> str:
    if codigo:
        return codigo
    slug = re.sub(r"[^A-Z0-9]", "", escola)[:24] or "ESCOLA"
    return f"CAD-{slug}"


def _nome_escola_preferido(atual: str, novo: str) -> str:
    if "RURAL" in novo.upper() and "RURAL)" in novo and "RURAL)" not in atual:
        return novo
    if len(novo) > len(atual):
        return novo
    return atual


def salvar_planilha_atual(origem: Path) -> Path:
    destino = DATA_DIR / "planilha-atual.xlsx"
    destino.write_bytes(origem.read_bytes())
    return destino


def completar_planilha(caminho: Path) -> dict:
    wb = load_workbook(caminho, data_only=True)
    from .homologacao import eh_planilha_caed, eh_planilha_homologacao, importar_caed, importar_homologacao

    if eh_planilha_caed(wb):
        with get_db() as conn:
            return importar_caed(caminho, conn)
    if not eh_planilha_homologacao(wb):
        raise ValueError("A planilha salva não é a de homologação de turmas.")
    with get_db() as conn:
        return importar_homologacao(caminho, conn, so_faltantes=True)


def importar_planilha(caminho: Path) -> dict:
    wb = load_workbook(caminho, data_only=True)
    from .homologacao import eh_planilha_caed, eh_planilha_homologacao, importar_caed, importar_homologacao

    if eh_planilha_caed(wb):
        with get_db() as conn:
            return importar_caed(caminho, conn)

    if eh_planilha_homologacao(wb):
        with get_db() as conn:
            return importar_homologacao(caminho, conn)

    ws = wb.active
    header_row, cols = _mapear_colunas(ws)

    avisos = []
    lidas = 0
    gravadas = 0
    viagem_por_mun: dict[str, list] = defaultdict(list)
    ordem_por_chave: dict[tuple, int] = defaultdict(int)

    with get_db() as conn:
        conn.execute("DELETE FROM vagas WHERE origem_linha IS NOT NULL")
        for row in ws.iter_rows(min_row=header_row + 1, max_col=16, values_only=False):
            cells_todos = [c.value for c in row]
            if all(v is None or str(v).strip() == "" for v in cells_todos):
                continue
            lidas += 1
            mun = _cab(_valor(row, cols["mun"]))
            escola = _cab(_valor(row, cols["escola"]))
            codigo = _codigo_escola(normalizar_texto(_valor(row, cols["codigo"])), escola)
            serie = normalizar_serie(_valor(row, cols["serie"]))
            turno = normalizar_turno(_valor(row, cols["turno"]))
            data = parse_data(_valor(row, cols["data"]))
            aplicador = normalizar_texto(_valor(row, cols["aplicador"]))
            saida = parse_data(_valor(row, cols["saida"]))
            retorno = parse_data(_valor(row, cols["retorno"]))

            if not mun or not escola or not serie or not data:
                avisos.append(f"Linha {row[0].row}: dados incompletos, ignorada.")
                continue

            municipio_id = garantir_municipio(conn, mun)

            rede = inferir_rede(escola)
            rural = 1 if eh_rural(escola) else 0
            existente = conn.execute(
                "SELECT id, nome FROM escolas WHERE codigo = ?", (codigo,)
            ).fetchone()
            if existente:
                escola_id = existente[0]
                nome_final = _nome_escola_preferido(existente[1], escola)
                conn.execute(
                    "UPDATE escolas SET nome=?, municipio_id=?, rede=?, rural=? WHERE id=?",
                    (nome_final, municipio_id, rede, rural, escola_id),
                )
            else:
                cur = conn.execute(
                    """INSERT INTO escolas(codigo, nome, municipio_id, rede, rural)
                       VALUES (?, ?, ?, ?, ?)""",
                    (codigo, escola, municipio_id, rede, rural),
                )
                escola_id = cur.lastrowid

            aplicador_id = None
            if aplicador:
                conn.execute(
                    "INSERT OR IGNORE INTO aplicadores(codigo, nome) VALUES (?, ?)",
                    (aplicador, aplicador),
                )
                n_apl = None
                m_apl = re.search(r"(\d+)", aplicador)
                if m_apl:
                    n_apl = int(m_apl.group(1))
                    conn.execute(
                        "UPDATE aplicadores SET numero = ? WHERE codigo = ?",
                        (n_apl, aplicador),
                    )
                aplicador_id = conn.execute(
                    "SELECT id FROM aplicadores WHERE codigo = ?", (aplicador,)
                ).fetchone()[0]

            chave_ordem = (escola_id, serie, turno)
            ordem_por_chave[chave_ordem] += 1
            ordem = ordem_por_chave[chave_ordem]
            conn.execute(
                """INSERT INTO vagas(escola_id, serie, turno, data, ordem, origem_linha, aplicador_id, status)
                   VALUES (?, ?, ?, ?, ?, ?, ?, 'PREVISTA')""",
                (escola_id, serie, turno, iso(data), ordem, row[0].row, aplicador_id),
            )
            gravadas += 1
            if saida or retorno:
                viagem_por_mun[mun].append((iso(saida), iso(retorno)))

        for mun, pares in viagem_por_mun.items():
            saidas = [p[0] for p in pares if p[0]]
            retornos = [p[1] for p in pares if p[1]]
            data_saida = min(saidas) if saidas else None
            data_retorno = max(retornos) if retornos else None
            mid = id_municipio_por_nome(conn, mun)
            if not mid:
                continue
            conn.execute(
                """INSERT INTO viagens(municipio_id, data_saida, data_retorno)
                   VALUES (?, ?, ?)
                   ON CONFLICT(municipio_id) DO UPDATE SET
                     data_saida=excluded.data_saida,
                     data_retorno=excluded.data_retorno""",
                (mid, data_saida, data_retorno),
            )

        consolidar_municipios(conn)

    return {
        "arquivo": caminho.name,
        "linhas_lidas": lidas,
        "vagas_gravadas": gravadas,
        "avisos": avisos[:30],
    }
