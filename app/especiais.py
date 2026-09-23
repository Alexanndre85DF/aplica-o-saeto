from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .homologacao import serie_da_etapa
from .regras import normalizar_texto, serie_base


def eh_planilha_confirmacao(wb) -> bool:
    return any("NECESSIDADE ESPECIAL" in str(n).upper() for n in wb.sheetnames)


def _cab(valor) -> str:
    return normalizar_texto(valor).upper()


def _flag(valor) -> str:
    texto = _cab(valor)
    if texto in {"SIM", "S"}:
        return "SIM"
    if texto in {"NÃO", "NAO", "N"}:
        return "NAO"
    return "NAO_INFORMADO"


def _aba_especiais(wb):
    for nome in wb.sheetnames:
        if "NECESSIDADE ESPECIAL" in str(nome).upper():
            return wb[nome]
    return None


def _vagas_do_aluno(conn, aluno) -> list[dict]:
    rows = conn.execute(
        """SELECT v.id, v.serie, v.turma, v.escola_id, v.n_extras
           FROM vagas v
           WHERE v.escola_id = ?""",
        (aluno["escola_id"],),
    ).fetchall()
    turma = normalizar_texto(aluno.get("turma"))
    base = serie_base(aluno.get("serie"))
    return [
        dict(r)
        for r in rows
        if normalizar_texto(r["turma"]) == turma and serie_base(r["serie"]) == base
    ]


def n_extras_da_vaga(conn, vaga) -> int:
    return sum(1 for a in alunos_da_vaga(conn, vaga) if a.get("precisa_extra"))


def _sincronizar_n_extras(conn, aluno=None, vaga=None) -> int:
    alvos = [vaga] if vaga else _vagas_do_aluno(conn, aluno or {})
    n = 0
    for item in alvos:
        n = n_extras_da_vaga(conn, item)
        conn.execute("UPDATE vagas SET n_extras = ? WHERE id = ?", (n, item["id"]))
    return n


def alinhar_n_extras_todas(conn) -> int:
    rows = conn.execute("SELECT id, escola_id, turma, serie FROM vagas").fetchall()
    for row in rows:
        _sincronizar_n_extras(conn, vaga=dict(row))
    return len(rows)


def alunos_da_vaga(conn, vaga) -> list[dict]:
    if not vaga or not vaga.get("escola_id"):
        return []
    try:
        rows = conn.execute(
            """SELECT * FROM alunos_especiais
               WHERE escola_id = ?
               ORDER BY nome""",
            (vaga["escola_id"],),
        ).fetchall()
    except Exception:
        return []
    turma = normalizar_texto(vaga.get("turma"))
    base = serie_base(vaga.get("serie"))
    vaga_id = vaga.get("id")
    extras = {}
    if vaga_id:
        try:
            for r in conn.execute(
                """SELECT x.aluno_id, a.id, a.codigo, a.nome
                   FROM vaga_extras x
                   JOIN aplicadores a ON a.id = x.aplicador_id
                   WHERE x.vaga_id = ? AND x.aluno_id IS NOT NULL""",
                (vaga_id,),
            ):
                extras[int(r["aluno_id"])] = {
                    "id": r["id"],
                    "codigo": r["codigo"],
                    "nome": r["nome"] or r["codigo"],
                }
        except Exception:
            extras = {}
    lista = []
    for r in rows:
        if normalizar_texto(r["turma"]) != turma:
            continue
        if serie_base(r["serie"]) != base:
            continue
        item = dict(r)
        item["precisa_extra"] = bool(item.get("precisa_extra"))
        item["extra"] = extras.get(int(item["id"]))
        lista.append(item)
    return lista


def criar_aluno_na_vaga(conn, vaga_id: int, nome: str, necessidade: str | None = None) -> dict:
    vaga = conn.execute(
        "SELECT id, escola_id, turma, serie, n_extras FROM vagas WHERE id = ?",
        (vaga_id,),
    ).fetchone()
    if not vaga:
        raise LookupError("Turma não encontrada.")
    nome = normalizar_texto(nome)
    if not nome:
        raise ValueError("Informe o nome do estudante.")
    turma = normalizar_texto(vaga["turma"])
    if not turma:
        raise ValueError("Esta vaga ainda não tem o nome da turma. Cadastre a turma antes.")
    serie = serie_base(vaga["serie"])
    conn.execute(
        """INSERT INTO alunos_especiais(
               escola_id, turma, serie, nome, necessidade, precisa_extra
           ) VALUES (?, ?, ?, ?, ?, 1)""",
        (vaga["escola_id"], turma, serie, nome, normalizar_texto(necessidade) or None),
    )
    aluno = {"escola_id": vaga["escola_id"], "turma": turma, "serie": serie}
    _sincronizar_n_extras(conn, aluno)
    return {"ok": True, "alunos": alunos_da_vaga(conn, dict(vaga))}


def importar_confirmacao(caminho: Path, conn) -> dict:
    wb = load_workbook(caminho, data_only=True)
    if not eh_planilha_confirmacao(wb):
        raise ValueError("Este arquivo não é o relatório de confirmação da base.")
    ws = _aba_especiais(wb)
    lidas = 0
    gravados = 0
    atualizados = 0
    sem_escola = 0
    avisos = []
    for row in ws.iter_rows(min_row=6, values_only=True):
        if not row or not row[13]:
            continue
        lidas += 1
        codigo_esc = str(row[4] or "").strip()
        turma = normalizar_texto(row[9])
        etapa = str(row[7] or "")
        serie = serie_base(serie_da_etapa(etapa) or serie_da_etapa(turma) or etapa)
        codigo_alu = str(row[12] or "").strip()
        nome = normalizar_texto(row[13])
        if not codigo_esc or not nome or not turma or not serie:
            avisos.append(f"{nome or 'Aluno'}: faltou escola, turma ou série.")
            continue
        escola = conn.execute(
            "SELECT id FROM escolas WHERE codigo = ?", (codigo_esc,)
        ).fetchone()
        if not escola:
            sem_escola += 1
            continue
        faz = _flag(row[16] if len(row) > 16 else None)
        precisa = 0 if faz == "SIM" else 1
        atual = None
        if codigo_alu:
            atual = conn.execute(
                "SELECT id FROM alunos_especiais WHERE codigo_estudante = ?",
                (codigo_alu,),
            ).fetchone()
        dados = (
            escola["id"],
            turma,
            serie,
            codigo_alu or None,
            nome,
            normalizar_texto(row[14]),
            normalizar_texto(row[15]),
            faz,
            _flag(row[19] if len(row) > 19 else None),
            _flag(row[18] if len(row) > 18 else None),
            precisa,
        )
        if atual:
            conn.execute(
                """UPDATE alunos_especiais
                   SET escola_id=?, turma=?, serie=?, codigo_estudante=?, nome=?,
                       necessidade=?, recurso=?, faz_com_turma=?, sala_extra=?,
                       profissional_escola=?, precisa_extra=?
                   WHERE id=?""",
                (*dados, atual["id"]),
            )
            aluno = {"escola_id": escola["id"], "turma": turma, "serie": serie}
            _sincronizar_n_extras(conn, aluno)
            atualizados += 1
        else:
            conn.execute(
                """INSERT INTO alunos_especiais(
                       escola_id, turma, serie, codigo_estudante, nome, necessidade,
                       recurso, faz_com_turma, sala_extra, profissional_escola, precisa_extra
                   ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                dados,
            )
            aluno = {"escola_id": escola["id"], "turma": turma, "serie": serie}
            _sincronizar_n_extras(conn, aluno)
            gravados += 1
    alinhar_n_extras_todas(conn)
    return {
        "ok": True,
        "tipo": "especiais",
        "arquivo": caminho.name,
        "linhas_lidas": lidas,
        "alunos": gravados + atualizados,
        "novos": gravados,
        "atualizados": atualizados,
        "sem_escola": sem_escola,
        "avisos": avisos[:40],
    }
