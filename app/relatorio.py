from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .regras import fmt_data

COLS = [
    (48, "Município"),
    (64, "Escola"),
    (18, "Rede"),
    (22, "Turno"),
    (22, "Data"),
    (34, "Série / turma"),
    (42, "Aplicador"),
    (14, "Extra"),
    (13, "Situação"),
]


_TROCAS = str.maketrans(
    {
        "\u2014": "-",
        "\u2013": "-",
        "\u2212": "-",
        "\u2022": "-",
        "\u2026": ".",
        "\u201c": '"',
        "\u201d": '"',
        "\u2018": "'",
        "\u2019": "'",
    }
)


def _txt(valor) -> str:
    texto = str(valor or "").translate(_TROCAS)
    return texto.encode("latin-1", "replace").decode("latin-1")


def _caber(pdf: FPDF, texto, largura: float) -> str:
    texto = _txt(texto)
    limite = max(6, largura - 1.6)
    if pdf.get_string_width(texto) <= limite:
        return texto
    while texto and pdf.get_string_width(texto + "...") > limite:
        texto = texto[:-1]
    return f"{texto}..." if texto else ""


def _celula(pdf: FPDF, largura: float, altura: float, texto, fill: bool = False, align: str = "L") -> None:
    pdf.cell(largura, altura, _caber(pdf, texto, largura), border=1, fill=fill, align=align)


def _situacao(slot: dict) -> str:
    if slot.get("status") == "FINALIZADA":
        return "Aplicada"
    if slot.get("vago"):
        return "Vago"
    if slot.get("tem_choque"):
        return "Choque"
    return "Alocado"


def _passa_filtro(linha: dict, slot: dict, rede: str, so_vagos: bool, status: str) -> bool:
    if rede and rede != "TODAS" and linha.get("rede") != rede:
        return False
    if so_vagos:
        return bool(slot.get("vago"))
    if status == "APLICADAS":
        return slot.get("status") == "FINALIZADA"
    if status == "PENDENTES":
        return (not slot.get("vago")) and slot.get("status") != "FINALIZADA"
    return True


def linhas_do_quadro(payload: dict, rede: str = "TODAS", so_vagos: bool = False, status: str = "TODAS") -> list[dict]:
    blocos = payload.get("quadros") if payload.get("todos") else [payload]
    saida = []
    for bloco in blocos or []:
        mun = ((bloco.get("municipio") or {}).get("nome") or "").strip()
        datas = bloco.get("datas") or []
        fmt = bloco.get("datas_fmt") or []
        rotulo = {d: f for d, f in zip(datas, fmt)}
        for linha in bloco.get("linhas") or []:
            for data in datas:
                for slot in (linha.get("celulas") or {}).get(data) or []:
                    if not _passa_filtro(linha, slot, rede, so_vagos, status):
                        continue
                    apl = slot.get("aplicador") or {}
                    n_ex = int(slot.get("n_extras") or 0)
                    serie = slot.get("serie") or ""
                    turma = slot.get("turma") or ""
                    saida.append(
                        {
                            "municipio": mun,
                            "escola": linha.get("escola") or "",
                            "rede": linha.get("rede") or "",
                            "turno": linha.get("turno") or "",
                            "data": rotulo.get(data) or (fmt_data(data) if data else "Sem data"),
                            "data_ord": data or "9999-99-99",
                            "serie": f"{serie} - {turma}".strip(" -") if turma else serie,
                            "aplicador": apl.get("nome") or apl.get("codigo") or "Sem aplicador",
                            "extra": f"{slot.get('extras_preenchidos') or 0}/{n_ex}" if n_ex else "-",
                            "situacao": _situacao(slot),
                        }
                    )
    saida.sort(key=lambda r: (r["municipio"], r["escola"], r["data_ord"], r["turno"], r["serie"]))
    return saida


def _rotulo_filtros(municipio_nome: str, rede: str, so_vagos: bool, status: str, total: int) -> str:
    partes = [f"Município: {municipio_nome or 'Todos'}"]
    partes.append(f"Rede: {rede.title() if rede and rede != 'TODAS' else 'Todas'}")
    if so_vagos:
        partes.append("Só vagos")
    else:
        sit = {
            "TODAS": "Todas as situações",
            "PENDENTES": "Ainda no campo",
            "APLICADAS": "Já aplicadas",
        }.get(status or "TODAS", status)
        partes.append(f"Situação: {sit}")
    partes.append(f"{total} aplicação(ões)")
    return "  -  ".join(partes)


class _PdfQuadro(FPDF):
    def __init__(self, filtros: str):
        super().__init__(orientation="L", unit="mm", format="A4")
        self.filtros = filtros
        self.set_margins(10, 12, 10)
        self.set_auto_page_break(auto=True, margin=14)

    def header(self):
        self.set_font("Helvetica", "B", 12)
        self.cell(0, 6, _txt("SAETO SRE Gurupi - Quadro de aplicação"), ln=True)
        self.set_font("Helvetica", "", 8)
        self.cell(0, 5, _txt(self.filtros), ln=True)
        self.ln(1)
        self.set_font("Helvetica", "B", 8)
        self.set_fill_color(0, 0, 128)
        self.set_text_color(255, 255, 255)
        for largura, titulo in COLS:
            self.cell(largura, 6, _txt(titulo), border=1, fill=True)
        self.ln()
        self.set_text_color(0, 0, 0)

    def footer(self):
        self.set_y(-10)
        self.set_font("Helvetica", "", 8)
        agora = datetime.now().strftime("%d/%m/%Y %H:%M")
        self.cell(0, 8, _txt(f"Emitido em {agora}  -  Página {self.page_no()}/{{nb}}"), align="C")


def gerar_pdf_quadro(
    payload: dict,
    *,
    municipio_nome: str,
    rede: str = "TODAS",
    so_vagos: bool = False,
    status: str = "TODAS",
) -> BytesIO:
    linhas = linhas_do_quadro(payload, rede, so_vagos, status)
    pdf = _PdfQuadro(_rotulo_filtros(municipio_nome, rede, so_vagos, status, len(linhas)))
    pdf.alias_nb_pages()
    pdf.add_page()
    pdf.set_font("Helvetica", "", 8)
    if not linhas:
        pdf.cell(0, 8, _txt("Nenhuma aplicação neste filtro."), ln=True)
    zebra = False
    for row in linhas:
        pdf.set_fill_color(236, 236, 245) if zebra else pdf.set_fill_color(255, 255, 255)
        valores = [
            row["municipio"],
            row["escola"],
            row["rede"],
            row["turno"],
            row["data"],
            row["serie"],
            row["aplicador"],
            row["extra"],
            row["situacao"],
        ]
        for largura, valor in zip((c[0] for c in COLS), valores):
            _celula(pdf, largura, 6, valor, fill=True)
        pdf.ln()
        zebra = not zebra
    bruto = pdf.output()
    return BytesIO(bytes(bruto))


def nome_arquivo_quadro(municipio_nome: str, ext: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", (municipio_nome or "todos")).strip("-").lower()[:40]
    return f"quadro-{slug or 'todos'}.{ext}"


def nome_arquivo_pdf(municipio_nome: str) -> str:
    return nome_arquivo_quadro(municipio_nome, "pdf")


def nome_arquivo_xlsx(municipio_nome: str) -> str:
    return nome_arquivo_quadro(municipio_nome, "xlsx")


def gerar_xlsx_quadro(
    payload: dict,
    *,
    municipio_nome: str,
    rede: str = "TODAS",
    so_vagos: bool = False,
    status: str = "TODAS",
) -> BytesIO:
    linhas = linhas_do_quadro(payload, rede, so_vagos, status)
    wb = Workbook()
    ws = wb.active
    ws.title = "Quadro"
    cab = [titulo for _, titulo in COLS]
    borda = Border(
        left=Side(style="thin", color="808080"),
        right=Side(style="thin", color="808080"),
        top=Side(style="thin", color="808080"),
        bottom=Side(style="thin", color="808080"),
    )
    fundo_cab = PatternFill("solid", fgColor="000080")
    fundo_zebra = PatternFill("solid", fgColor="ECECF5")
    fonte_cab = Font(bold=True, color="FFFFFF")
    ws.merge_cells("A1:I1")
    ws["A1"] = "SAETO SRE Gurupi — Quadro de aplicação"
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A2:I2")
    ws["A2"] = _rotulo_filtros(municipio_nome, rede, so_vagos, status, len(linhas))
    ws.append([])
    ws.append(cab)
    for cell in ws[4]:
        cell.font = fonte_cab
        cell.fill = fundo_cab
        cell.border = borda
        cell.alignment = Alignment(horizontal="center")
    for i, row in enumerate(linhas):
        ws.append(
            [
                row["municipio"],
                row["escola"],
                row["rede"],
                row["turno"],
                row["data"],
                row["serie"],
                row["aplicador"],
                row["extra"],
                row["situacao"],
            ]
        )
        for cell in ws[ws.max_row]:
            cell.border = borda
            if i % 2:
                cell.fill = fundo_zebra
    if not linhas:
        ws.append(["Nenhuma aplicação neste filtro."])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 28
    ws.column_dimensions["G"].width = 32
    ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 12
    ws.auto_filter.ref = f"A4:I{max(4, ws.max_row)}"
    ws.freeze_panes = "A5"
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


COLS_DIARIA = [
    (32, "Data"),
    (78, "Município"),
    (88, "Servidor"),
    (30, "Matrícula"),
    (22, "OS"),
    (27, "Valor"),
]


def _linha_diaria_passa(linha: dict, q: str) -> bool:
    if linha.get("excluido"):
        return False
    if not q:
        return True
    bloco = " ".join(
        [
            str(linha.get("nome") or ""),
            str(linha.get("matricula") or ""),
            str(linha.get("rota") or ""),
            str(linha.get("os") or ""),
            str(linha.get("municipio") or ""),
        ]
    ).lower()
    return q in bloco


def grupos_diaria_filtrados(folha: dict, municipio_id: int | None = None, q: str | None = None) -> list[dict]:
    busca = (q or "").strip().lower()
    saida = []
    for grupo in folha.get("grupos") or []:
        if municipio_id and int(grupo.get("municipio_id") or 0) != int(municipio_id):
            continue
        linhas = [l for l in (grupo.get("linhas") or []) if _linha_diaria_passa(l, busca)]
        if not linhas:
            continue
        total = sum(float(l.get("valor") or 0) for l in linhas)
        saida.append({**grupo, "linhas": linhas, "total": total})
    return saida


class _PdfDiaria(FPDF):
    def __init__(self, filtros: str):
        super().__init__(orientation="L", unit="mm", format="A4")
        self.filtros = filtros
        self.set_margins(10, 12, 10)
        self.set_auto_page_break(auto=True, margin=14)

    def header(self):
        self.set_font("Helvetica", "B", 12)
        self.cell(0, 6, _txt("SAETO SRE Gurupi - Folha de diárias"), ln=True)
        self.set_font("Helvetica", "", 8)
        self.cell(0, 5, _txt(self.filtros), ln=True)
        self.ln(1)
        self.set_font("Helvetica", "B", 8)
        self.set_fill_color(0, 0, 128)
        self.set_text_color(255, 255, 255)
        for largura, titulo in COLS_DIARIA:
            self.cell(largura, 6, _txt(titulo), border=1, fill=True)
        self.ln()
        self.set_text_color(0, 0, 0)

    def footer(self):
        self.set_y(-10)
        self.set_font("Helvetica", "", 8)
        agora = datetime.now().strftime("%d/%m/%Y %H:%M")
        self.cell(0, 8, _txt(f"Emitido em {agora}  -  Página {self.page_no()}/{{nb}}"), align="C")


def _rotulo_diarias(folha: dict, grupos: list[dict], municipio_id: int | None, q: str | None) -> str:
    from .diarias import fmt_moeda

    pessoas = {l.get("aplicador_id") for g in grupos for l in g["linhas"]}
    total = sum(g["total"] for g in grupos)
    destino = "Todos os destinos"
    if municipio_id:
        for g in grupos:
            destino = g.get("municipio") or destino
            break
        if destino == "Todos os destinos":
            for g in folha.get("grupos") or []:
                if int(g.get("municipio_id") or 0) == int(municipio_id):
                    destino = g.get("municipio") or destino
                    break
    partes = [f"Destino: {destino}"]
    if (q or "").strip():
        partes.append(f"Busca: {(q or '').strip()}")
    partes.append(f"{len(pessoas)} pessoa(s)")
    partes.append(f"Total: {fmt_moeda(total)}")
    return "  -  ".join(partes)


def gerar_pdf_diarias(folha: dict, municipio_id: int | None = None, q: str | None = None) -> BytesIO:
    from .diarias import fmt_moeda

    grupos = grupos_diaria_filtrados(folha, municipio_id, q)
    total = sum(g["total"] for g in grupos)
    pdf = _PdfDiaria(_rotulo_diarias(folha, grupos, municipio_id, q))
    pdf.alias_nb_pages()
    pdf.add_page()
    if not grupos:
        pdf.set_font("Helvetica", "", 8)
        pdf.cell(0, 8, _txt("Nenhuma diária neste filtro."), ln=True)
    for grupo in grupos:
        pdf.set_font("Helvetica", "", 8)
        zebra = False
        for linha in grupo["linhas"]:
            pdf.set_fill_color(236, 236, 245) if zebra else pdf.set_fill_color(255, 255, 255)
            valores = [
                linha.get("data_fmt") or "-",
                linha.get("rota") or linha.get("municipio") or "",
                (linha.get("nome") or "").upper(),
                linha.get("matricula") or "-",
                linha.get("os") or "",
                linha.get("valor_fmt") or "",
            ]
            for largura, valor in zip((c[0] for c in COLS_DIARIA), valores):
                _celula(pdf, largura, 6, valor, fill=True)
            pdf.ln()
            zebra = not zebra
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_fill_color(255, 242, 204)
        largura_rotulo = sum(c[0] for c in COLS_DIARIA[:-1])
        _celula(pdf, largura_rotulo, 6, f"Total {grupo.get('data_fmt') or ''} - {grupo.get('rota') or ''}", fill=True)
        _celula(pdf, COLS_DIARIA[-1][0], 6, fmt_moeda(grupo["total"]), fill=True, align="R")
        pdf.ln()
    if grupos:
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_fill_color(217, 226, 243)
        largura_rotulo = sum(c[0] for c in COLS_DIARIA[:-1])
        _celula(pdf, largura_rotulo, 7, "Total geral", fill=True)
        _celula(pdf, COLS_DIARIA[-1][0], 7, fmt_moeda(total), fill=True, align="R")
        pdf.ln()
    return BytesIO(bytes(pdf.output()))


def gerar_xlsx_diarias(folha: dict, municipio_id: int | None = None, q: str | None = None) -> BytesIO:
    from .diarias import fmt_moeda

    grupos = grupos_diaria_filtrados(folha, municipio_id, q)
    total = sum(g["total"] for g in grupos)
    wb = Workbook()
    ws = wb.active
    ws.title = "Diárias"
    cab = ["DATA", "MUNICÍPIO", "SERVIDOR", "MATRÍCULA", "OS", "VALOR"]
    borda = Border(
        left=Side(style="thin", color="808080"),
        right=Side(style="thin", color="808080"),
        top=Side(style="thin", color="808080"),
        bottom=Side(style="thin", color="808080"),
    )
    fundo_cab = PatternFill("solid", fgColor="000080")
    fundo_tot = PatternFill("solid", fgColor="FFF2CC")
    fundo_geral = PatternFill("solid", fgColor="D9E2F3")
    fonte_cab = Font(bold=True, color="FFFFFF")
    negrito = Font(bold=True)
    ws.merge_cells("A1:F1")
    ws["A1"] = "SAETO SRE Gurupi — Folha de diárias"
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A2:F2")
    ws["A2"] = _rotulo_diarias(folha, grupos, municipio_id, q)
    ws.append([])
    ws.append(cab)
    for cell in ws[4]:
        cell.font = fonte_cab
        cell.fill = fundo_cab
        cell.border = borda
        cell.alignment = Alignment(horizontal="center")
    if not grupos:
        ws.append(["Nenhuma diária neste filtro."])
    for grupo in grupos:
        for linha in grupo["linhas"]:
            ws.append(
                [
                    linha.get("data_fmt") or "—",
                    linha.get("rota") or linha.get("municipio") or "",
                    (linha.get("nome") or "").upper(),
                    linha.get("matricula") or "",
                    linha.get("os") or "",
                    linha.get("valor_fmt") or "",
                ]
            )
            for cell in ws[ws.max_row]:
                cell.border = borda
        ws.append(
            [
                "",
                "",
                f"Total {grupo.get('data_fmt') or ''} · {grupo.get('rota') or ''}",
                "",
                "",
                fmt_moeda(grupo["total"]),
            ]
        )
        for cell in ws[ws.max_row]:
            cell.font = negrito
            cell.fill = fundo_tot
            cell.border = borda
    if grupos:
        ws.append(["", "", "Total geral", "", "", fmt_moeda(total)])
        for cell in ws[ws.max_row]:
            cell.font = negrito
            cell.fill = fundo_geral
            cell.border = borda
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 16
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
